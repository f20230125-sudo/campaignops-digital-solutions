import csv,json,unittest
from pathlib import Path
import pymupdf
from openpyxl import load_workbook
ROOT=Path(__file__).resolve().parents[1]
class ReportReconciliation(unittest.TestCase):
    def test_downloads_match_dashboard_and_manifest(self):
        data=json.loads((ROOT/'public/data/dashboard.json').read_text(encoding='utf-8'));total=next(m for m in data['metrics'] if m['campaign']==m['market']==m['period']=='all')
        with (ROOT/'public/downloads/campaign-metrics.csv').open(newline='') as f:rows=list(csv.DictReader(f))
        row=next(m for m in rows if m['campaign']==m['market']==m['period']=='all')
        for key in ['delivered','sessions','registrations','attendees']:self.assertEqual(int(row[key]),total[key])
        wb=load_workbook(ROOT/'public/downloads/campaign-metrics.xlsx',data_only=True);self.assertEqual(wb['Scoped metrics'].max_row,len(data['metrics'])+1)
        self.assertIn(data['manifest']['data_hash'],[r[1] for r in wb['Provenance'].values]);wb.close()
        doc=pymupdf.open(ROOT/'public/downloads/campaign-report.pdf');text=''.join(p.get_text() for p in doc);doc.close();self.assertIn(data['manifest']['data_hash'],text)
        for key in data['registry']:self.assertIn(f"{total[key]:.1%}",text)
        manifest=json.loads((ROOT/'public/downloads/report-manifest.json').read_text());self.assertEqual(manifest['data_hash'],data['manifest']['data_hash'])
