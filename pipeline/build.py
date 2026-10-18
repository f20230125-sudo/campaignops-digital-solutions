"""Validate immutable exports, calculate scoped SQL metrics and publish traceable reports."""
import argparse, csv, hashlib, json, re, tempfile
from collections import Counter
from datetime import date, datetime, timedelta, timezone
from pathlib import Path
import duckdb
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
import pymupdf
from .generate import ROOT, FIELDS

VERSION='1.0.0'
CUTOFF=datetime(2026,10,1,tzinfo=timezone.utc)
KINDS={'delivered','click','session','registration','webinar_registration','attendance'}
REGISTRY={
 'email_click_rate': {'label':'Email click rate','numerator':'clicked','denominator':'delivered','definition':'Delivered messages with at least one qualifying human click / delivered messages. Repeat clicks collapse to message ID; known bot clicks excluded.'},
 'engagement_rate': {'label':'Engaged sessions','numerator':'engaged','denominator':'sessions','definition':'Eligible landing sessions lasting more than 10 seconds, with at least two pageviews, or with a registration / eligible landing sessions.'},
 'registration_rate': {'label':'Registration rate','numerator':'registrations','denominator':'sessions','definition':'Eligible landing sessions with a registration / eligible landing sessions.'},
 'attendance_rate': {'label':'Webinar attendance','numerator':'attendees','denominator':'registrants','definition':'Distinct registered webinar/person pairs with positive attendee time / distinct webinar/person registrations. Rejoins collapse; presenters excluded. Periods use event date.'},
}

def digest(data): return hashlib.sha256(data).hexdigest()
def write_json(path, value):
    path.parent.mkdir(parents=True,exist_ok=True); path.write_text(json.dumps(value,indent=2,ensure_ascii=False)+'\n',encoding='utf-8')

def validate_files(paths, mappings, cutoff=CUTOFF, provenance='synthetic'):
    candidates={}; rejects=[]; imports=[]; hashes=set(); duplicate_events=0
    for path in sorted(paths):
        filehash=digest(path.read_bytes()); duplicate_file=filehash in hashes; hashes.add(filehash)
        imports.append(dict(file=path.name,sha256=filehash,duplicate=duplicate_file))
        if duplicate_file: continue
        with path.open(newline='',encoding='utf-8-sig') as f:
            reader=csv.DictReader(f)
            if reader.fieldnames!=FIELDS: raise ValueError(f'{path.name}: expected version 1 columns {FIELDS}')
            for line,row in enumerate(reader,2):
                reason=None
                try:
                    timestamp=datetime.fromisoformat(row['timestamp'].replace('Z','+00:00'))
                    if timestamp.utcoffset() is None: raise ValueError('naive timestamp')
                    timestamp=timestamp.astimezone(timezone.utc)
                except (ValueError,TypeError): timestamp=None; reason='invalid_timestamp'
                if not reason and timestamp>=cutoff: reason='after_cutoff'
                if row['campaign_raw'] not in mappings: reason=reason or 'unmapped_campaign'
                if row['market'] not in {'AE','SA','QA'}: reason=reason or 'invalid_market'
                if row['kind'] not in KINDS: reason=reason or 'invalid_kind'
                if row['provenance']!=provenance: reason=reason or 'wrong_provenance'
                if not row['event_id'] or not row['entity_id']: reason=reason or 'missing_id'
                if row['kind'] in {'session','attendance'}:
                    try:
                        number=float(row['value'])
                        if not 0<=number<=86400: raise ValueError()
                    except (ValueError,TypeError): reason=reason or 'invalid_value'
                if row['kind']=='session' and not re.fullmatch(r'[1-9][0-9]*\|landing\|(mobile|desktop)',row['flag']): reason=reason or 'invalid_session_flag'
                if row['kind']=='click' and row['flag'] not in {'human','bot'}: reason=reason or 'invalid_click_flag'
                if row['kind']=='attendance' and (row['flag'] not in {'attendee','presenter','organizer'} or not row['parent_id']): reason=reason or 'invalid_attendance'
                if row['kind']=='webinar_registration' and not row['parent_id']: reason=reason or 'missing_webinar'
                rec=dict(source_file=path.name,line=line,reason=reason,raw=row)
                if reason: rejects.append(rec); continue
                identity=row['event_id']; payload=json.dumps(row,sort_keys=True)
                candidates.setdefault(identity,[]).append((payload,rec,timestamp))
    accepted=[]
    for identity, group in candidates.items():
        if len({g[0] for g in group})>1:
            for _,rec,_ in group: rec['reason']='conflicting_event_id'; rejects.append(rec)
            continue
        duplicate_events+=len(group)-1
        _,rec,timestamp=group[0]; row=rec['raw']
        accepted.append(dict(event_id=identity,kind=row['kind'],campaign=mappings[row['campaign_raw']],campaign_raw=row['campaign_raw'],market=row['market'],occurred_at=timestamp.isoformat(),entity_id=row['entity_id'],parent_id=row['parent_id'],value=float(row['value'] or 0),flag=row['flag'],provenance=provenance,source_file=rec['source_file'],source_line=rec['line']))
    # An orphan registration is evidence of a missing session, not a conversion.
    session_ids={(r['campaign'],r['market'],r['entity_id']) for r in accepted if r['kind']=='session'}
    orphan_count=sum(r['kind']=='registration' and (r['campaign'],r['market'],r['entity_id']) not in session_ids for r in accepted)
    return accepted,rejects,dict(imports=imports,duplicate_events=duplicate_events,orphan_registrations=orphan_count)

def periods():
    result=[dict(id='all',label='Full period',start='2026-04-01',end='2026-09-30',type='all')]
    for month in range(4,10):
        start=date(2026,month,1); end=date(2026,month+1,1)-timedelta(days=1)
        result.append(dict(id=start.strftime('%Y-%m'),label=start.strftime('%B 2026'),start=str(start),end=str(end),type='month'))
    start=date(2026,3,30)
    while start<=date(2026,9,30):
        result.append(dict(id=f'week-{start}',label=f'Week of {start:%d %b}',start=str(start),end=str(start+timedelta(days=6)),type='week')); start+=timedelta(days=7)
    return result

def calculate(accepted, ps, connection=None):
    con=connection or duckdb.connect()
    con.execute('CREATE OR REPLACE TABLE events (event_id VARCHAR, kind VARCHAR,campaign VARCHAR,campaign_raw VARCHAR,market VARCHAR,occurred_at TIMESTAMPTZ,entity_id VARCHAR,parent_id VARCHAR,value DOUBLE,flag VARCHAR,provenance VARCHAR,source_file VARCHAR,source_line INTEGER)')
    if accepted:
        with tempfile.NamedTemporaryFile(mode='w', suffix='.json', delete=False, encoding='utf-8') as staged:
            json.dump(accepted, staged); staged_path=Path(staged.name)
        try:
            con.read_json(str(staged_path), format='array', columns={k:'VARCHAR' for k in accepted[0]}).create_view('incoming_events')
            con.execute('INSERT INTO events SELECT * FROM incoming_events')
            con.execute('DROP VIEW incoming_events')
        finally: staged_path.unlink()
    con.execute('CREATE OR REPLACE TABLE periods (id VARCHAR,start_date DATE,end_date DATE)')
    con.executemany('INSERT INTO periods VALUES (?,?,?)',[(p['id'],p['start'],p['end']) for p in ps])
    cursor=con.execute((ROOT/'pipeline/metrics.sql').read_text(encoding='utf-8')); names=[d[0] for d in cursor.description]
    result=[dict(zip(names,row)) for row in cursor.fetchall()]
    for row in result:
        for key,metric in REGISTRY.items():
            denominator=row[metric['denominator']]; row[key]=row[metric['numerator']]/denominator if denominator else None
    return result

def build():
    maps=json.loads((ROOT/'data/reference/mappings.json').read_text(encoding='utf-8')); campaigns=json.loads((ROOT/'data/reference/campaigns.json').read_text(encoding='utf-8')); assets=json.loads((ROOT/'data/reference/assets.json').read_text(encoding='utf-8'))
    accepted,rejects,audit=validate_files((ROOT/'data/raw').glob('*.csv'),maps)
    calculations=((ROOT/'pipeline/metrics.sql').read_text(encoding='utf-8').replace('\r\n','\n')+(ROOT/'pipeline/build.py').read_text(encoding='utf-8').replace('\r\n','\n')).encode()
    datahash=digest(json.dumps(dict(imports=audit['imports'],mappings=maps,calculation_hash=digest(calculations),cutoff=CUTOFF.isoformat()),sort_keys=True).encode())
    out=ROOT/'public/data'; out.mkdir(parents=True,exist_ok=True)
    runtime=ROOT/'runtime'; runtime.mkdir(exist_ok=True)
    con=duckdb.connect(str(runtime/'campaignops.duckdb')); ps=periods(); metrics=calculate(accepted,ps,con); con.close()
    quality=Counter(r['reason'] for r in rejects)
    manifest=dict(version=VERSION,contract_version='events-v1',data_hash=datahash,calculation_sha256=digest(calculations),cutoff=CUTOFF.isoformat(),report_timezone='Asia/Dubai',provenance='synthetic',accepted=len(accepted),rejected=len(rejects),quality=dict(quality),**audit)
    issues=[
        dict(id='utm',title='Campaign tags split the hepatitis journey',severity='High',campaign='hepatitis',category='Tracking',count=quality.get('unmapped_campaign',0),evidence='Hepatitis_2026 appears in web exports but has no approved mapping.',action='Review and add the explicit mapping, then rebuild. Preserve the original raw tag.',resolution='Reclassification can recover attributable records; it cannot create new audience activity.'),
        dict(id='duplicate',title='A webinar file was exported twice',severity='Medium',campaign='all',category='Import',count=sum(i['duplicate'] for i in audit['imports']),evidence='webinar.csv and webinar-reexport.csv have identical SHA-256 hashes. Attendee rejoins also exist.',action='Skip the identical file. Count webinar/person pairs, not attendance rows.',resolution='Exact file deduplication is automatic; conflicting event IDs are quarantined.'),
        dict(id='mobile',title='The hepatitis registration link is hidden on mobile',severity='High',campaign='hepatitis',category='Website',count=1,evidence='Seeded WordPress landing pages hide their registration action at narrow widths.',action='Edit the asset metadata in WordPress, publish and run a fresh mobile test journey.',resolution='A passing test verifies the current path. Historical conversion rates stay unchanged.'),
    ]
    write_json(out/'dashboard.json',dict(campaigns=campaigns,assets=assets,periods=ps,metrics=metrics,registry=REGISTRY,manifest=manifest,issues=issues))
    write_json(out/'quality.json',dict(manifest=manifest,rejects=rejects))
    for campaign in campaigns: write_json(out/f"records-{campaign['id']}.json",[r for r in accepted if r['campaign']==campaign['id']])
    write_json(out/'manifest.json',manifest)
    exports=ROOT/'public/downloads'; exports.mkdir(parents=True,exist_ok=True)
    columns=list(metrics[0]);
    with (exports/'campaign-metrics.csv').open('w',newline='',encoding='utf-8') as f:
        writer=csv.DictWriter(f,fieldnames=columns); writer.writeheader(); writer.writerows(metrics)
    workbook=Workbook(); sheet=workbook.active; sheet.title='Scoped metrics'; sheet.append(columns)
    for row in metrics: sheet.append([row[k] for k in columns])
    sheet.freeze_panes='D2'; sheet.auto_filter.ref=sheet.dimensions
    for cell in sheet[1]: cell.font=Font(color='FFFFFF',bold=True); cell.fill=PatternFill('solid',fgColor='183E35')
    for idx,key in enumerate(columns,1):
        sheet.column_dimensions[sheet.cell(1,idx).column_letter].width=22
        if key in REGISTRY:
            for col in sheet.iter_cols(min_col=idx,max_col=idx,min_row=2):
                for cell in col: cell.number_format='0.0%'
    provenance=workbook.create_sheet('Provenance'); provenance.append(['Field','Value'])
    for k in ['data_hash','calculation_sha256','cutoff','version','provenance','report_timezone']: provenance.append([k,str(manifest[k])])
    defs=workbook.create_sheet('Metric definitions'); defs.append(['Metric','Definition'])
    for k,v in REGISTRY.items(): defs.append([k,v['definition']])
    workbook.save(exports/'campaign-metrics.xlsx')
    # Dependency-light, searchable PDF. Scope, denominator and quality travel with the report.
    doc=pymupdf.open(); page=doc.new_page(width=595,height=842)
    def text(y,s,size=11): page.insert_text((44,y),s,fontsize=size,fontname='helv',color=(.08,.19,.16))
    total=next(r for r in metrics if r['period']==r['campaign']==r['market']=='all')
    text(52,'CAMPAIGNOPS  /  PERFORMANCE REPORT',18); text(81,'April-September 2026 | All campaigns | AE, SA, QA')
    text(105,'Synthetic demonstration activity. Not MSD or real audience performance.',10)
    y=150
    for key,m in REGISTRY.items():
        text(y,f"{m['label']}: {total[key]:.1%} ({total[m['numerator']]:,} / {total[m['denominator']]:,})",14); y+=40
    text(330,f"Data accepted: {len(accepted):,} events | Quarantined: {len(rejects)}",11)
    text(360,'Quality: unmapped tags, a conflicting event ID, invalid time and late feed.',10)
    text(390,'Duplicate webinar file skipped by hash. Rejoins collapse to participant/webinar.',10)
    text(430,'Cutoff: '+manifest['cutoff'],10); text(450,'Report timezone: Asia/Dubai | Calculation version: '+VERSION,10)
    text(490,'Input manifest SHA-256:',10); text(509,datahash,8)
    text(539,'Calculation SHA-256:',10); text(558,manifest['calculation_sha256'],8)
    text(600,'Definitions and all weekly/monthly scopes: campaign-metrics.xlsx',10)
    text(625,'No ROI, clinical recommendation or historical conversion-lift claim.',10)
    doc.save(exports/'campaign-report.pdf'); doc.close()
    write_json(exports/'report-manifest.json',manifest)
    from .weekly import build_weekly
    build_weekly()
    print(json.dumps(dict(accepted=len(accepted),rejected=len(rejects),snapshots=len(metrics),quality=dict(quality),data_hash=datahash),indent=2))

if __name__=='__main__':
    parser=argparse.ArgumentParser(); parser.parse_args(); build()
